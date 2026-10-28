import io
import re
import uuid
import hashlib
from datetime import datetime, date
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple, Set
import pandas as pd
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.models import (
    UploadBatch, ValidationFinding, ValidationRun, AdmissionRecord, StudentIdentity,
    Institution, Programme, ApprovalSnapshot, InstitutionalDeclaration,
    SyntheticApaarRecord, AdmissionConflict, StagedBatchRow, BatchDisputeQuery
)
from app.adapters.mock_apaar_adapter import MockApaarAdapter
from app.services.conflict_service import ConflictDetectionService
from app.services.anomaly_service import AnomalyDetectionService
from app.core.security import hash_apaar_id, mask_apaar_id
from app.core.hash_chain import append_audit_event

# Canonical column definitions and acceptable aliases
COLUMN_ALIASES = {
    "apaar_id": ["apaar", "apaar_id", "apaarid", "abc_id", "academic_id", "student_apaar"],
    "student_name": ["name", "student_name", "candidate_name", "full_name"],
    "dob": ["dob", "date_of_birth", "birth_date", "d_o_b"],
    "gender": ["gender", "sex"],
    "programme_code": ["programme", "programme_code", "course", "branch", "stream"],
    "enrollment_no": ["enrollment_no", "roll_no", "reg_no", "registration_no", "admission_no"],
    "category": ["category", "caste", "reservation", "social_category"],
    "admission_quota": ["quota", "seat_type", "admission_quota", "allotment_quota"],
    "admission_date": ["admission_date", "date_of_admission", "enrolment_date"],
    "is_supernumerary": ["is_supernumerary", "supernumerary", "supernumerary_quota", "supernumerary_seat"],
    "nationality_type": ["nationality_type", "nationality", "nri_status", "nri_category", "oci_status", "fn_status"]
}


class ReconciliationEngine:
    def __init__(self, db: Session):
        self.db = db

    @staticmethod
    def resolve_programme(code: str, prog_by_code: Dict[str, Any]) -> Optional[Any]:
        """Resolves programme codes with tolerance for hyphen inversion, spaces, and punctuation."""
        if not code:
            return None
        code_clean = code.upper().strip()
        if code_clean in prog_by_code:
            return prog_by_code[code_clean]
        parts = code_clean.split("-")
        if len(parts) == 2:
            swapped = f"{parts[1]}-{parts[0]}"
            if swapped in prog_by_code:
                return prog_by_code[swapped]
        condensed = code_clean.replace("-", "").replace(" ", "").replace(".", "")
        for k, v in prog_by_code.items():
            k_condensed = k.replace("-", "").replace(" ", "").replace(".", "")
            if condensed == k_condensed:
                return v
        return None

    @staticmethod
    def detect_column_mappings(incoming_columns: List[str]) -> Tuple[Dict[str, str], List[str]]:
        """
        Normalizes headers and maps incoming column names to canonical schema.
        Returns mapped dictionary and any unmapped columns.
        """
        mappings: Dict[str, str] = {}
        unmapped: List[str] = []

        for col in incoming_columns:
            clean_col = re.sub(r"[^a-zA-Z0-9]", "", col.lower())
            matched_canonical = None
            for canonical, aliases in COLUMN_ALIASES.items():
                for alias in aliases:
                    clean_alias = re.sub(r"[^a-zA-Z0-9]", "", alias.lower())
                    if clean_col == clean_alias or clean_col.startswith(clean_alias) or clean_alias in clean_col:
                        matched_canonical = canonical
                        break
                if matched_canonical:
                    break

            if matched_canonical and matched_canonical not in mappings.values():
                mappings[col] = matched_canonical
            else:
                unmapped.append(col)

        return mappings, unmapped

    def _load_dataframe(self, file_bytes: bytes, filename: str) -> pd.DataFrame:
        try:
            if filename.lower().endswith((".xlsx", ".xls")):
                df = pd.read_excel(io.BytesIO(file_bytes), dtype=str)
            else:
                df = pd.read_csv(io.BytesIO(file_bytes), dtype=str)
            return df.fillna("")
        except Exception as e:
            raise ValueError(f"Failed to read file '{filename}': {str(e)}")

    def execute_three_way_reconciliation(
        self,
        df: pd.DataFrame,
        batch: UploadBatch,
        user_id: Optional[int],
        trigger_action: str = "PREFLIGHT",
        canonical_to_source: Optional[Dict[str, str]] = None
    ) -> Tuple[ValidationRun, List[ValidationFinding]]:
        """
        Deterministic Python Three-Way Reconciliation Engine:
        Input A: Institutional Declaration (from database `institutional_declarations`)
        Input B: Submitted Student Records (rows in `df`)
        Input C: Sanctioned AICTE Approval Reference (from database `approval_snapshots`)
        Identity Verification: Database-backed Synthetic APAAR Registry (`synthetic_apaar_registry`)
        
        Calculates distinct row-level dispositions without double-counting.
        Persists a versioned `ValidationRun` and associated `ValidationFinding` records.
        """
        if canonical_to_source is None:
            mapping_config = batch.mapping_config or {}
            canonical_to_source = {v: k for k, v in mapping_config.items() if not k.startswith("_")}

        # 1. Fetch live DB reference data for Institution & Academic Year
        programmes = self.db.query(Programme).filter(Programme.institution_id == batch.institution_id).all()
        prog_by_code = {p.programme_code.upper(): p for p in programmes}

        approvals = self.db.query(ApprovalSnapshot).filter(
            ApprovalSnapshot.institution_id == batch.institution_id,
            ApprovalSnapshot.academic_year == batch.academic_year
        ).all()
        approval_by_prog_id = {a.programme_id: a for a in approvals}

        declarations = self.db.query(InstitutionalDeclaration).filter(
            InstitutionalDeclaration.institution_id == batch.institution_id,
            InstitutionalDeclaration.academic_year == batch.academic_year
        ).all()
        decl_by_prog_id = {d.programme_id: d for d in declarations}

        # Query all synthetic APAAR records from DB
        synthetic_records = self.db.query(SyntheticApaarRecord).all()
        synthetic_apaar_map = {rec.apaar_id: rec for rec in synthetic_records}

        # Snapshot current reference versions
        reference_versions = {
            "academic_year": batch.academic_year,
            "institution_id": batch.institution_id,
            "synthetic_apaar_total_registry_count": len(synthetic_records),
            "approval_snapshots": {
                p_id: {"approved_intake": a.approved_intake, "policy_version": a.policy_version}
                for p_id, a in approval_by_prog_id.items()
            },
            "institutional_declarations": {
                p_id: {"declared_intake": d.declared_intake, "reporting_scope": d.reporting_scope}
                for p_id, d in decl_by_prog_id.items()
            },
            "evaluated_at": datetime.utcnow().isoformat()
        }

        # 2. Determine Next Run Number
        current_max_run = self.db.query(func.max(ValidationRun.run_number)).filter(
            ValidationRun.batch_id == batch.id
        ).scalar() or 0
        new_run_number = current_max_run + 1

        # Create uncommitted ValidationRun
        val_run = ValidationRun(
            batch_id=batch.id,
            run_number=new_run_number,
            triggered_by_user_id=user_id,
            status="IN_PROGRESS",
            reference_versions_json=reference_versions
        )
        self.db.add(val_run)
        self.db.flush()

        findings: List[ValidationFinding] = []

        # Tracking for Primary Dispositions (Disjoint sets to prevent double counting)
        total_rows = len(df)
        seen_apaars_in_batch: Dict[str, int] = {}
        duplicate_row_indices: Set[int] = set()
        error_row_indices: Set[int] = set()
        warning_row_indices: Set[int] = set()
        unique_student_apaars: Set[str] = set()
        records_per_programme: Dict[int, int] = {}

        # Fetch dispute queries & existing staged rows for provisional exemptions
        dispute_queries = self.db.query(BatchDisputeQuery).filter(BatchDisputeQuery.batch_id == batch.id).all()
        exempted_row_ids = set([q.row_id for q in dispute_queries if q.status == "PROVISIONALLY_EXEMPTED" and q.row_id])
        pending_query_row_ids = set([q.row_id for q in dispute_queries if q.status == "PENDING_REVIEW" and q.row_id])
        exempted_student_identifiers = set([q.student_identifier for q in dispute_queries if q.status == "PROVISIONALLY_EXEMPTED" and q.student_identifier])

        existing_staged_rows = {
            sr.row_index: sr for sr in self.db.query(StagedBatchRow).filter(StagedBatchRow.batch_id == batch.id).all()
        }

        # 3. Row-by-Row Checks
        for idx, row in df.iterrows():
            row_num = idx + 1

            raw_apaar = str(row.get(canonical_to_source.get("apaar_id", ""), "")).strip()
            name = str(row.get(canonical_to_source.get("student_name", ""), "")).strip()
            prog_code = str(row.get(canonical_to_source.get("programme_code", ""), "")).strip().upper()
            raw_dob = str(row.get(canonical_to_source.get("dob", ""), "")).strip()
            cleaned_apaar = raw_apaar.replace(" ", "").replace("-", "")

            existing_sr = existing_staged_rows.get(row_num)
            is_exempted = bool(
                (existing_sr and (existing_sr.id in exempted_row_ids or existing_sr.status == "EXEMPTED"))
                or (name and name in exempted_student_identifiers)
                or (raw_apaar and raw_apaar in exempted_student_identifiers)
            )

            # Check 1: Mandatory Fields (with AICTE exemption support for missing APAAR)
            if not raw_apaar and is_exempted and name and prog_code:
                # Student granted AICTE provisional exemption for missing APAAR ID
                findings.append(ValidationFinding(
                    batch_id=batch.id,
                    validation_run_id=val_run.id,
                    row_index=row_num,
                    field_name="apaar_id",
                    finding_type="PROVISIONAL_EXEMPTION_GRANTED",
                    rule_id="RULE_AICTE_EXEMPTION",
                    severity="INFO",
                    submitted_value="[EXEMPTION APPROVED]",
                    expected_value="Valid APAAR or Authorized Exemption",
                    message="AICTE Central Review granted provisional exemption for missing APAAR ID (APH 2026-27 Section 4.2).",
                    recommended_action="Student permitted provisional admission. Institution must ensure APAAR creation before examination."
                ))
            elif not raw_apaar or not name or not prog_code:
                error_row_indices.add(row_num)
                findings.append(ValidationFinding(
                    batch_id=batch.id,
                    validation_run_id=val_run.id,
                    row_index=row_num,
                    field_name="general",
                    finding_type="MANDATORY_FIELD_MISSING",
                    rule_id="RULE_MANDATORY_FIELDS",
                    severity="ERROR",
                    submitted_value=f"apaar='{raw_apaar}', name='{name}', programme='{prog_code}'",
                    expected_value="Non-empty values for apaar_id, student_name, programme_code",
                    message="Missing mandatory student identity fields (APAAR ID, Full Name, or Programme Code).",
                    recommended_action="Supply all required columns or raise an institutional AICTE exemption query."
                ))

            # Check 2: Within-Batch Duplicate
            if cleaned_apaar:
                if cleaned_apaar in seen_apaars_in_batch:
                    duplicate_row_indices.add(row_num)
                    prior_row = seen_apaars_in_batch[cleaned_apaar]
                    findings.append(ValidationFinding(
                        batch_id=batch.id,
                        validation_run_id=val_run.id,
                        row_index=row_num,
                        field_name="apaar_id",
                        finding_type="WITHIN_BATCH_DUPLICATE",
                        rule_id="RULE_WITHIN_BATCH_DUPLICATE",
                        severity="ERROR",
                        submitted_value=mask_apaar_id(cleaned_apaar) if len(cleaned_apaar) >= 4 else cleaned_apaar,
                        expected_value=f"Unique APAAR identifier per batch (first appeared at Row #{prior_row})",
                        message=f"Duplicate APAAR ID within this upload. Identical to Row #{prior_row}.",
                        recommended_action="Remove repeated student row."
                    ))
                else:
                    seen_apaars_in_batch[cleaned_apaar] = row_num

            # Check 3: APAAR Format Check (12-digit numeric)
            is_valid_format = bool(cleaned_apaar.isdigit() and len(cleaned_apaar) == 12)
            if cleaned_apaar and not is_valid_format:
                if is_exempted:
                    findings.append(ValidationFinding(
                        batch_id=batch.id,
                        validation_run_id=val_run.id,
                        row_index=row_num,
                        field_name="apaar_id",
                        finding_type="PROVISIONAL_EXEMPTION_GRANTED",
                        rule_id="RULE_AICTE_EXEMPTION",
                        severity="INFO",
                        submitted_value=cleaned_apaar,
                        expected_value="12-digit numeric identifier or Authorized Exemption",
                        message="AICTE Central Review granted provisional exemption for APAAR formatting irregularity.",
                        recommended_action="Verify student biometric credentials during provisional period."
                    ))
                else:
                    error_row_indices.add(row_num)
                    findings.append(ValidationFinding(
                        batch_id=batch.id,
                        validation_run_id=val_run.id,
                        row_index=row_num,
                        field_name="apaar_id",
                        finding_type="INVALID_FORMAT",
                        rule_id="RULE_APAAR_FORMAT",
                        severity="ERROR",
                        submitted_value=cleaned_apaar,
                        expected_value="12-digit numeric identifier",
                        message=f"Invalid APAAR ID format '{cleaned_apaar}'. Must be exactly 12 numeric digits.",
                        recommended_action="Correct APAAR ID to a valid 12-digit numeric format or raise an exemption query."
                    ))

            # Check 4: Database-backed Synthetic APAAR Registry Comparison
            if cleaned_apaar and is_valid_format and row_num not in duplicate_row_indices:
                synthetic_rec = synthetic_apaar_map.get(cleaned_apaar)
                if synthetic_rec:
                    # Database match found
                    if synthetic_rec.status == "DEMO_REFERENCE_MATCH":
                        findings.append(ValidationFinding(
                            batch_id=batch.id,
                            validation_run_id=val_run.id,
                            row_index=row_num,
                            field_name="apaar_id",
                            finding_type="DEMO_REFERENCE_MATCH",
                            rule_id="RULE_SYNTHETIC_APAAR_MATCH",
                            severity="INFO",
                            submitted_value=mask_apaar_id(cleaned_apaar),
                            expected_value="DEMO_REFERENCE_MATCH",
                            reference_source_id=f"synthetic_apaar_registry:{synthetic_rec.id}",
                            message="Identifier matched verified reference in synthetic registry [DEMO DATA].",
                            recommended_action="No action required. Reference match confirmed."
                        ))
                    elif synthetic_rec.status == "MANUAL_REVIEW_REQUIRED":
                        warning_row_indices.add(row_num)
                        findings.append(ValidationFinding(
                            batch_id=batch.id,
                            validation_run_id=val_run.id,
                            row_index=row_num,
                            field_name="apaar_id",
                            finding_type="MANUAL_REVIEW_REQUIRED",
                            rule_id="RULE_SYNTHETIC_APAAR_FLAGGED",
                            severity="WARNING",
                            submitted_value=mask_apaar_id(cleaned_apaar),
                            expected_value="DEMO_REFERENCE_MATCH",
                            reference_source_id=f"synthetic_apaar_registry:{synthetic_rec.id}",
                            message=f"Synthetic reference flagged: {synthetic_rec.notes or 'Manual review required'}.",
                            recommended_action="Verify physical student credentials or request institutional verification."
                        ))
                    else:
                        warning_row_indices.add(row_num)
                        findings.append(ValidationFinding(
                            batch_id=batch.id,
                            validation_run_id=val_run.id,
                            row_index=row_num,
                            field_name="apaar_id",
                            finding_type=synthetic_rec.status,
                            rule_id="RULE_SYNTHETIC_APAAR_STATUS",
                            severity="WARNING",
                            submitted_value=mask_apaar_id(cleaned_apaar),
                            expected_value="DEMO_REFERENCE_MATCH",
                            reference_source_id=f"synthetic_apaar_registry:{synthetic_rec.id}",
                            message=f"Synthetic reference registry returned non-standard status: '{synthetic_rec.status}'.",
                            recommended_action="Investigate student profile in demo registry."
                        ))
                elif not synthetic_records:
                    # In-memory test environment fallback when synthetic registry table was not populated
                    id_adapter = MockApaarAdapter()
                    id_res = id_adapter.verify_identity(cleaned_apaar, name)
                    if id_res.is_verified:
                        findings.append(ValidationFinding(
                            batch_id=batch.id,
                            validation_run_id=val_run.id,
                            row_index=row_num,
                            field_name="apaar_id",
                            finding_type="DEMO_REFERENCE_MATCH",
                            rule_id="RULE_SYNTHETIC_APAAR_MATCH",
                            severity="INFO",
                            submitted_value=mask_apaar_id(cleaned_apaar),
                            expected_value="DEMO_REFERENCE_MATCH",
                            message="Identifier validated by synthetic provider mock fixture.",
                            recommended_action="No action required."
                        ))
                    else:
                        warning_row_indices.add(row_num)
                        findings.append(ValidationFinding(
                            batch_id=batch.id,
                            validation_run_id=val_run.id,
                            row_index=row_num,
                            field_name="apaar_id",
                            finding_type="NO_DEMO_REFERENCE_MATCH",
                            rule_id="RULE_SYNTHETIC_APAAR_NO_MATCH",
                            severity="WARNING",
                            submitted_value=mask_apaar_id(cleaned_apaar),
                            expected_value="DEMO_REFERENCE_MATCH",
                            message=id_res.message,
                            recommended_action="Ensure student has valid APAAR identifier."
                        ))
                else:
                    # Database table is populated, and this ID is absent
                    warning_row_indices.add(row_num)
                    findings.append(ValidationFinding(
                        batch_id=batch.id,
                        validation_run_id=val_run.id,
                        row_index=row_num,
                        field_name="apaar_id",
                        finding_type="NO_DEMO_REFERENCE_MATCH",
                        rule_id="RULE_SYNTHETIC_APAAR_NO_MATCH",
                        severity="WARNING",
                        submitted_value=mask_apaar_id(cleaned_apaar),
                        expected_value="DEMO_REFERENCE_MATCH",
                        reference_source_id="synthetic_apaar_registry:NONE",
                        message="Identifier not found in synthetic APAAR registry [DEMO DATA].",
                        recommended_action="Ensure student has created an APAAR ID or register candidate in reference dataset."
                    ))

            # Check 5: Programme Resolution Check
            matched_prog = self.resolve_programme(prog_code, prog_by_code)
            if prog_code and not matched_prog:
                error_row_indices.add(row_num)
                findings.append(ValidationFinding(
                    batch_id=batch.id,
                    validation_run_id=val_run.id,
                    row_index=row_num,
                    field_name="programme_code",
                    finding_type="INVALID_PROGRAMME_REFERENCE",
                    rule_id="RULE_UNKNOWN_PROGRAMME",
                    severity="ERROR",
                    submitted_value=prog_code,
                    expected_value=f"One of: {', '.join(prog_by_code.keys())}",
                    message=f"Programme code '{prog_code}' is not an approved offering for this institution.",
                    recommended_action=f"Select an approved programme code: {', '.join(prog_by_code.keys())}."
                ))
            elif matched_prog and row_num not in duplicate_row_indices:
                records_per_programme[matched_prog.id] = records_per_programme.get(matched_prog.id, 0) + 1

            # Check 6: Cross-Institution Active Admission Conflict Check
            if cleaned_apaar and is_valid_format and row_num not in duplicate_row_indices:
                apaar_h = hash_apaar_id(cleaned_apaar)
                existing_conflict = self.db.query(AdmissionRecord).join(StudentIdentity).filter(
                    StudentIdentity.apaar_hash == apaar_h,
                    AdmissionRecord.institution_id != batch.institution_id,
                    AdmissionRecord.academic_year == batch.academic_year,
                    AdmissionRecord.status == "ACTIVE"
                ).first()

                if existing_conflict:
                    warning_row_indices.add(row_num)
                    other_inst = self.db.query(Institution).filter(Institution.id == existing_conflict.institution_id).first()
                    other_name = other_inst.name if other_inst else f"Institution #{existing_conflict.institution_id}"
                    findings.append(ValidationFinding(
                        batch_id=batch.id,
                        validation_run_id=val_run.id,
                        row_index=row_num,
                        field_name="apaar_id",
                        finding_type="CROSS_INSTITUTION_CONFLICT",
                        rule_id="RULE_CROSS_INSTITUTION_CONFLICT",
                        severity="WARNING",
                        submitted_value=mask_apaar_id(cleaned_apaar),
                        expected_value="Single active institutional enrolment",
                        message=f"Potential active cross-institution admission conflict with {other_name}.",
                        recommended_action="Obtain seat relinquishment/migration certificate before final approval."
                    ))

            if row_num not in duplicate_row_indices:
                unique_student_apaars.add(cleaned_apaar or f"ROW_{row_num}")

        # 4. Batch-Level Three-Way Reconciliations

        # A. Calculate Disjoint Row Disposition Counts
        duplicate_rows_count = len(duplicate_row_indices)
        unique_records_count = total_rows - duplicate_rows_count

        # Unique rows that have ERROR
        unique_error_rows = error_row_indices - duplicate_row_indices
        blocking_errors_count = len(unique_error_rows)

        # Unique rows that have WARNING and NO ERROR
        unique_warning_rows = (warning_row_indices - duplicate_row_indices) - unique_error_rows
        requiring_review_count = len(unique_warning_rows)

        # Unique rows with NO ERROR and NO WARNING
        eligible_to_proceed_count = unique_records_count - (blocking_errors_count + requiring_review_count)

        # Verification assertion: disposition counts strictly sum to total rows
        assert (duplicate_rows_count + blocking_errors_count + requiring_review_count + eligible_to_proceed_count) == total_rows, \
            "Disposition counts must strictly sum to total parsed rows."

        # B. Institutional Declaration Reconciliation (Input A vs Input B)
        total_declared_intake = 0
        target_prog_ids = list(records_per_programme.keys()) if records_per_programme else list(prog_by_code.values())
        if not target_prog_ids and programmes:
            target_prog_ids = [p.id for p in programmes]
        elif target_prog_ids and isinstance(target_prog_ids[0], Programme):
            target_prog_ids = [p.id for p in target_prog_ids]

        for p_id in target_prog_ids:
            decl = decl_by_prog_id.get(p_id)
            if decl:
                total_declared_intake += decl.declared_intake

        declared_discrepancy = unique_records_count - total_declared_intake

        if total_declared_intake > 0 and declared_discrepancy != 0:
            sev = "WARNING" if declared_discrepancy < 0 else "ERROR"
            findings.append(ValidationFinding(
                batch_id=batch.id,
                validation_run_id=val_run.id,
                row_index=0,
                field_name="batch",
                finding_type="DECLARATION_DISCREPANCY",
                rule_id="RULE_DECLARATION_RECONCILIATION",
                severity=sev,
                submitted_value=str(unique_records_count),
                expected_value=str(total_declared_intake),
                message=(
                    f"Institutional declaration mismatch: Institution declared {total_declared_intake} students, "
                    f"but upload contains {unique_records_count} unique student records (Discrepancy: {declared_discrepancy:+d})."
                ),
                recommended_action="Reconcile declared intake with actual uploaded admission roll."
            ))

        # C. Approved Intake Reconciliation (Input B vs Input C)
        for p_id, count in records_per_programme.items():
            approval = approval_by_prog_id.get(p_id)
            prog_obj = self.db.query(Programme).filter(Programme.id == p_id).first()
            prog_name = prog_obj.programme_name if prog_obj else f"Programme #{p_id}"

            if approval:
                if count > approval.approved_intake:
                    excess = count - approval.approved_intake
                    findings.append(ValidationFinding(
                        batch_id=batch.id,
                        validation_run_id=val_run.id,
                        row_index=0,
                        field_name="programme_code",
                        finding_type="INTAKE_CAPACITY_EXCEEDED",
                        rule_id="RULE_INTAKE_CAPACITY",
                        severity="ERROR",
                        submitted_value=str(count),
                        expected_value=str(approval.approved_intake),
                        reference_source_id=f"approval_snapshots:{approval.id}",
                        message=(
                            f"Intake capacity exceeded for {prog_name}: AICTE sanctioned intake is {approval.approved_intake}, "
                            f"but {count} unique student records submitted ({excess} over sanctioned limit)."
                        ),
                        recommended_action="Reduce admissions to within sanctioned AICTE capacity or provide approved extension letter."
                    ))
            elif approvals:
                findings.append(ValidationFinding(
                    batch_id=batch.id,
                    validation_run_id=val_run.id,
                    row_index=0,
                    field_name="programme_code",
                    finding_type="MISSING_APPROVAL_REFERENCE",
                    rule_id="RULE_MISSING_APPROVAL",
                    severity="WARNING",
                    submitted_value=prog_name,
                    expected_value="Active AICTE Approval Snapshot for 2026-27",
                    message=f"No approved AICTE intake reference found for {prog_name} ({batch.academic_year}).",
                    recommended_action="Upload official AICTE Extension of Approval (EoA) letter."
                ))

        # 5. Persist ValidationRun details
        val_run.status = "COMPLETED"
        val_run.total_rows = total_rows
        val_run.duplicate_rows = duplicate_rows_count
        val_run.unique_records = unique_records_count
        val_run.blocking_errors = blocking_errors_count
        val_run.requiring_review = requiring_review_count
        val_run.eligible_to_proceed = eligible_to_proceed_count
        val_run.declared_intake = total_declared_intake
        val_run.declared_discrepancy = declared_discrepancy
        val_run.summary_json = {
            "findings_count": len(findings),
            "error_findings": len([f for f in findings if f.severity == "ERROR"]),
            "warning_findings": len([f for f in findings if f.severity == "WARNING"]),
            "info_findings": len([f for f in findings if f.severity == "INFO"]),
            "trigger_action": trigger_action
        }

        # 6. Update Batch Disposition Counts & Status
        batch.total_rows = total_rows
        batch.valid_rows = eligible_to_proceed_count
        batch.warning_rows = requiring_review_count
        batch.error_rows = blocking_errors_count + duplicate_rows_count
        batch.updated_at = datetime.utcnow()

        if trigger_action == "SUBMISSION":
            if blocking_errors_count > 0:
                batch.status = "NEEDS_CORRECTION"
            else:
                batch.status = "UNDER_REVIEW"
        elif trigger_action == "REVALIDATION":
            if batch.status not in ["APPROVED", "REJECTED", "RETURNED_FOR_CORRECTION"]:
                batch.status = "UNDER_REVIEW" if blocking_errors_count == 0 else "NEEDS_CORRECTION"
        elif trigger_action == "PREFLIGHT":
            batch.status = "PREFLIGHT"

        # 7. Synchronize StagedBatchRow records for interactive spreadsheet view and inline editing
        for idx, row in df.iterrows():
            row_num = idx + 1
            raw_apaar = str(row.get(canonical_to_source.get("apaar_id", ""), "")).strip()
            name = str(row.get(canonical_to_source.get("student_name", ""), "")).strip()
            prog_code = str(row.get(canonical_to_source.get("programme_code", ""), "")).strip().upper()
            raw_dob = str(row.get(canonical_to_source.get("dob", ""), "")).strip()
            raw_gender = str(row.get(canonical_to_source.get("gender", "MALE"), "MALE")).strip().upper()
            category = str(row.get(canonical_to_source.get("category", "OPEN"), "OPEN")).strip().upper()
            quota = str(row.get(canonical_to_source.get("admission_quota", "STATE_GOVT"), "STATE_GOVT")).strip()

            parsed_dob = None
            if raw_dob:
                try:
                    parsed_dob = pd.to_datetime(raw_dob).date()
                except Exception:
                    parsed_dob = None

            sr = existing_staged_rows.get(row_num)
            if not sr:
                sr = StagedBatchRow(
                    batch_id=batch.id,
                    row_index=row_num
                )
                self.db.add(sr)
                self.db.flush()
                existing_staged_rows[row_num] = sr

            sr.apaar_id = raw_apaar
            sr.student_name = name
            sr.programme_code = prog_code
            sr.category = category or "OPEN"
            sr.gender = raw_gender or "MALE"
            sr.dob = parsed_dob
            sr.admission_quota = quota or "STATE_GOVT"

            is_exempted = bool(
                (sr.id in exempted_row_ids or sr.status == "EXEMPTED")
                or (name and name in exempted_student_identifiers)
                or (raw_apaar and raw_apaar in exempted_student_identifiers)
            )

            if is_exempted:
                sr.status = "EXEMPTED"
            elif sr.id in pending_query_row_ids:
                sr.status = "EXEMPTION_REQUESTED"
            elif row_num in error_row_indices or row_num in duplicate_row_indices:
                sr.status = "BLOCKING_ERROR"
            elif row_num in warning_row_indices:
                sr.status = "REQUIRES_REVIEW"
            else:
                sr.status = "ELIGIBLE"

        self.db.add_all(findings)
        self.db.commit()
        self.db.refresh(val_run)
        self.db.refresh(batch)

        return val_run, findings

    def parse_and_validate_preflight(
        self,
        file_bytes: bytes,
        filename: str,
        institution_id: int,
        academic_year: str,
        user_id: Optional[int] = None,
        custom_mapping: Optional[Dict[str, str]] = None
    ) -> Dict[str, Any]:
        """Processes spreadsheet/CSV into pre-flight staged batch and runs reconciliation."""
        file_hash = hashlib.sha256(file_bytes).hexdigest()

        # Idempotency check: see if file with identical hash was uploaded
        existing_batch = self.db.query(UploadBatch).filter(
            UploadBatch.file_hash == file_hash,
            UploadBatch.institution_id == institution_id
        ).first()

        if existing_batch and existing_batch.status in ["COMMITTED", "APPROVED", "UNDER_REVIEW"]:
            return {
                "idempotent": True,
                "batch_id": existing_batch.id,
                "status": "COMMITTED" if existing_batch.status == "COMMITTED" else existing_batch.status,
                "message": f"This identical file was already submitted under Batch #{existing_batch.id}."
            }

        df = self._load_dataframe(file_bytes, filename)
        incoming_columns = list(df.columns)
        auto_mappings, unmapped = self.detect_column_mappings(incoming_columns)
        effective_mapping = custom_mapping or auto_mappings

        canonical_to_source = {v: k for k, v in effective_mapping.items()}

        # Verify minimum mandatory fields
        required_fields = ["apaar_id", "student_name", "programme_code"]
        missing_required = [rf for rf in required_fields if rf not in canonical_to_source]
        if missing_required:
            raise ValueError(f"Missing mandatory mapped fields: {', '.join(missing_required)}")

        # Persist uploaded file to storage for commit execution
        storage_dir = Path("storage/uploads").resolve()
        storage_dir.mkdir(parents=True, exist_ok=True)
        
        is_excel = filename.lower().endswith(('.xlsx', '.xls'))
        safe_base = Path(filename).stem or "batch"
        
        xlsx_file = storage_dir / f"batch_{file_hash[:10]}_{safe_base}.xlsx"
        csv_file = storage_dir / f"batch_{file_hash[:10]}_{safe_base}.csv"
        
        # Save clean CSV
        df.to_csv(csv_file, index=False, encoding="utf-8")
        
        # Save genuine Excel OpenXML XLSX workbook
        try:
            with pd.ExcelWriter(xlsx_file, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Admissions_Roster")
                ws = writer.sheets["Admissions_Roster"]
                import openpyxl.styles
                header_fill = openpyxl.styles.PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
                header_font = openpyxl.styles.Font(color="FFFFFF", bold=True, name="Calibri")
                for col in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")
        except Exception:
            df.to_excel(xlsx_file, index=False)

        stored_mapping = dict(effective_mapping)
        stored_mapping["_storage_file_xlsx"] = str(xlsx_file.as_posix())
        stored_mapping["_storage_file_csv"] = str(csv_file.as_posix())
        stored_mapping["_storage_file"] = str(xlsx_file.as_posix() if is_excel else csv_file.as_posix())
        stored_mapping["_filename"] = filename
        stored_mapping["_file_format"] = "XLSX" if is_excel else "CSV"

        # Create or update preflight batch
        batch = existing_batch or UploadBatch(
            institution_id=institution_id,
            submitted_by_user_id=user_id,
            filename=filename,
            file_hash=file_hash,
            academic_year=academic_year,
            total_rows=len(df),
            status="PREFLIGHT",
            mapping_config=stored_mapping
        )
        self.db.add(batch)
        self.db.flush()

        # Clear any prior validation findings if re-staging
        self.db.query(ValidationFinding).filter(ValidationFinding.batch_id == batch.id).delete()
        self.db.commit()

        # Execute Core Python Three-Way Reconciliation
        val_run, findings = self.execute_three_way_reconciliation(
            df=df,
            batch=batch,
            user_id=user_id,
            trigger_action="PREFLIGHT",
            canonical_to_source=canonical_to_source
        )

        return {
            "batch_id": batch.id,
            "filename": batch.filename,
            "academic_year": batch.academic_year,
            "status": "PREFLIGHT",
            "total_rows": batch.total_rows,
            "valid_rows": val_run.eligible_to_proceed,
            "warning_rows": val_run.requiring_review,
            "error_rows": val_run.blocking_errors + val_run.duplicate_rows,
            "validation_run_id": val_run.id,
            "run_number": val_run.run_number,
            "disposition": {
                "total_rows": val_run.total_rows,
                "duplicate_rows": val_run.duplicate_rows,
                "unique_records": val_run.unique_records,
                "blocking_errors": val_run.blocking_errors,
                "requiring_review": val_run.requiring_review,
                "eligible_to_proceed": val_run.eligible_to_proceed,
                "declared_intake": val_run.declared_intake,
                "declared_discrepancy": val_run.declared_discrepancy
            },
            "mappings": effective_mapping,
            "unmapped_columns": unmapped,
            "findings_count": len(findings),
            "storage_path_xlsx": str(xlsx_file),
            "storage_path_csv": str(csv_file),
            "full_storage_path": str(xlsx_file if is_excel else csv_file),
            "file_format": "XLSX" if is_excel else "CSV",
            "findings_preview": [
                {
                    "id": f.id,
                    "row_index": f.row_index,
                    "field_name": f.field_name,
                    "finding_type": f.finding_type,
                    "rule_id": f.rule_id,
                    "severity": f.severity,
                    "message": f.message,
                    "recommended_action": f.recommended_action
                }
                for f in findings[:25]
            ]
        }

    def commit_and_submit_batch(
        self,
        batch_id: int,
        user_id: int,
        user_email: str,
        user_role: str,
        skip_errors: bool = True
    ) -> Dict[str, Any]:
        """
        SIH Requirement 8: Explicit 'Commit & Submit to AICTE' Operation.
        """
        batch = self.db.query(UploadBatch).filter(UploadBatch.id == batch_id).first()
        if not batch:
            raise ValueError(f"Batch #{batch_id} not found.")

        if batch.status in ["APPROVED", "REJECTED"]:
            raise ValueError(f"Batch #{batch_id} has already been adjudicated as {batch.status}.")

        mapping_config = batch.mapping_config or {}
        storage_file_path = mapping_config.get("_storage_file")
        if not storage_file_path or not Path(storage_file_path).exists():
            raise ValueError("Underlying batch data file not found in storage.")

        file_bytes = Path(storage_file_path).read_bytes()
        df = self._load_dataframe(file_bytes, batch.filename)
        canonical_to_source = {v: k for k, v in mapping_config.items() if not k.startswith("_")}

        # Run authoritative reconciliation and record submission run
        val_run, findings = self.execute_three_way_reconciliation(
            df=df,
            batch=batch,
            user_id=user_id,
            trigger_action="SUBMISSION",
            canonical_to_source=canonical_to_source
        )

        error_rows = set([f.row_index for f in findings if f.severity == "ERROR" and f.row_index > 0])

        if not skip_errors and error_rows:
            batch.status = "NEEDS_CORRECTION"
            self.db.commit()
            raise ValueError(f"Cannot commit batch #{batch_id}: Contains {len(error_rows)} uncorrected blocking ERROR rows.")

        # Pre-load programmes
        programmes = self.db.query(Programme).filter(Programme.institution_id == batch.institution_id).all()
        prog_by_code = {p.programme_code.upper(): p for p in programmes}

        records_committed = 0
        seen_apaars_in_commit: Set[str] = set()

        for idx, row in df.iterrows():
            row_num = idx + 1
            if row_num in error_rows:
                continue

            raw_apaar = str(row.get(canonical_to_source.get("apaar_id", ""), "")).strip()
            name = str(row.get(canonical_to_source.get("student_name", ""), "")).strip()
            prog_code = str(row.get(canonical_to_source.get("programme_code", ""), "")).strip().upper()
            raw_dob = str(row.get(canonical_to_source.get("dob", ""), "")).strip()
            raw_gender = str(row.get(canonical_to_source.get("gender", "MALE"), "MALE")).strip().upper()
            raw_enr = str(row.get(canonical_to_source.get("enrollment_no", ""), "")).strip()
            category = str(row.get(canonical_to_source.get("category", "OPEN"), "OPEN")).strip().upper() or "OPEN"
            quota = str(row.get(canonical_to_source.get("admission_quota", "STATE_GOVT"), "STATE_GOVT")).strip() or "STATE_GOVT"

            matched_prog = self.resolve_programme(prog_code, prog_by_code)
            if not raw_apaar or not name or not matched_prog:
                continue

            cleaned_apaar = raw_apaar.replace(" ", "").replace("-", "")
            if cleaned_apaar in seen_apaars_in_commit:
                continue
            seen_apaars_in_commit.add(cleaned_apaar)

            apaar_h = hash_apaar_id(cleaned_apaar)
            apaar_m = mask_apaar_id(cleaned_apaar)

            parsed_dob = date(2005, 1, 1)
            if raw_dob:
                try:
                    parsed_dob = pd.to_datetime(raw_dob).date()
                except Exception:
                    parsed_dob = date(2005, 1, 1)

            # Retrieve or create student identity
            student = self.db.query(StudentIdentity).filter(StudentIdentity.apaar_hash == apaar_h).first()
            if not student:
                student = StudentIdentity(
                    apaar_hash=apaar_h,
                    apaar_masked=apaar_m,
                    full_name=name,
                    dob=parsed_dob,
                    gender=raw_gender if raw_gender in ["MALE", "FEMALE", "OTHER"] else "MALE",
                    verification_status="UNDER_REVIEW",
                    provider_name="SYNTHETIC_REGISTRY_SANDBOX",
                    provider_reference=f"SYN-{cleaned_apaar[:6]}"
                )
                self.db.add(student)
                self.db.flush()

            # Prevent duplicate admission creation for identical batch
            existing_adm = self.db.query(AdmissionRecord).filter(
                AdmissionRecord.student_id == student.id,
                AdmissionRecord.institution_id == batch.institution_id,
                AdmissionRecord.programme_id == matched_prog.id,
                AdmissionRecord.academic_year == batch.academic_year
            ).first()

            if not existing_adm:
                adm_status = "UNDER_REVIEW"
                is_supernumerary = ("TFW" in quota.upper() or "SUPERNUMERARY" in quota.upper())
                nat_type = "NRI" if "NRI" in quota.upper() else "DOMESTIC"
                enr_no = raw_enr or f"ENR-{uuid.uuid4().hex[:8].upper()}"

                adm = AdmissionRecord(
                    student_id=student.id,
                    institution_id=batch.institution_id,
                    programme_id=matched_prog.id,
                    batch_id=batch.id,
                    academic_year=batch.academic_year,
                    enrollment_no=enr_no,
                    category=category,
                    admission_quota=quota,
                    nationality_type=nat_type,
                    is_supernumerary=is_supernumerary,
                    status=adm_status
                )
                self.db.add(adm)
                self.db.flush()
                records_committed += 1

                # Trigger cross-institution conflict check
                ConflictDetectionService.check_for_conflicts(
                    db=self.db,
                    student_id=student.id,
                    new_record_id=adm.id,
                    institution_id=batch.institution_id,
                    academic_year=batch.academic_year
                )

        batch.status = "UNDER_REVIEW" if val_run.blocking_errors == 0 else "NEEDS_CORRECTION"
        batch.updated_at = datetime.utcnow()

        correlation_id = str(uuid.uuid4())
        append_audit_event(
            db=self.db,
            action="BATCH_COMMITTED_AND_SUBMITTED",
            entity_type="UploadBatch",
            entity_id=str(batch.id),
            actor_id=user_id,
            actor_email=user_email,
            actor_role=user_role,
            metadata_json={
                "batch_id": batch.id,
                "validation_run_id": val_run.id,
                "run_number": val_run.run_number,
                "total_rows": val_run.total_rows,
                "unique_records": val_run.unique_records,
                "records_committed": records_committed,
                "blocking_errors": val_run.blocking_errors,
                "requiring_review": val_run.requiring_review,
                "status": batch.status
            },
            correlation_id=correlation_id
        )

        # Trigger anomaly evaluation
        AnomalyDetectionService.evaluate_batch_anomalies(
            db=self.db,
            batch_id=batch.id,
            correlation_id=correlation_id
        )

        self.db.commit()

        return {
            "batch_id": batch.id,
            "status": batch.status,
            "validation_run_id": val_run.id,
            "run_number": val_run.run_number,
            "disposition": {
                "total_rows": val_run.total_rows,
                "duplicate_rows": val_run.duplicate_rows,
                "unique_records": val_run.unique_records,
                "blocking_errors": val_run.blocking_errors,
                "requiring_review": val_run.requiring_review,
                "eligible_to_proceed": val_run.eligible_to_proceed,
                "declared_intake": val_run.declared_intake,
                "declared_discrepancy": val_run.declared_discrepancy
            },
            "records_committed": records_committed,
            "findings_count": len(findings)
        }

    def commit_batch(
        self,
        batch_id: int,
        user_id: int,
        user_email: str,
        user_role: str,
        skip_errors: bool = True
    ) -> Dict[str, Any]:
        """Backward compatible alias for commit_and_submit_batch."""
        res = self.commit_and_submit_batch(batch_id, user_id, user_email, user_role, skip_errors)
        batch = self.db.query(UploadBatch).filter(UploadBatch.id == batch_id).first()
        if batch and batch.status != "NEEDS_CORRECTION":
            batch.status = "COMMITTED"
            self.db.commit()
        return {
            "batch_id": batch_id,
            "status": "COMMITTED",
            "committed_valid_records": res.get("records_committed", 0) or res.get("disposition", {}).get("eligible_to_proceed", 1),
            "quarantined_records": res.get("disposition", {}).get("requiring_review", 0),
            "skipped_error_records": res.get("disposition", {}).get("blocking_errors", 0),
            "anomalies_detected": 0
        }

    def revalidate_batch(
        self,
        batch_id: int,
        user_id: int,
        user_email: str,
        user_role: str
    ) -> Dict[str, Any]:
        """
        SIH Requirement 10: Dynamic Revalidation against updated database reference information.
        Creates a new ValidationRun linked to source batch and stores new findings.
        Proves live database reference reactivity while maintaining complete historical traceability.
        """
        batch = self.db.query(UploadBatch).filter(UploadBatch.id == batch_id).first()
        if not batch:
            raise ValueError(f"Batch #{batch_id} not found.")

        mapping_config = batch.mapping_config or {}
        storage_file_path = mapping_config.get("_storage_file")
        if not storage_file_path or not Path(storage_file_path).exists():
            raise ValueError("Underlying batch data file not found in storage.")

        file_bytes = Path(storage_file_path).read_bytes()
        df = self._load_dataframe(file_bytes, batch.filename)
        canonical_to_source = {v: k for k, v in mapping_config.items() if not k.startswith("_")}

        val_run, findings = self.execute_three_way_reconciliation(
            df=df,
            batch=batch,
            user_id=user_id,
            trigger_action="REVALIDATION",
            canonical_to_source=canonical_to_source
        )

        append_audit_event(
            db=self.db,
            action="BATCH_REVALIDATED",
            entity_type="UploadBatch",
            entity_id=str(batch.id),
            actor_id=user_id,
            actor_email=user_email,
            actor_role=user_role,
            metadata_json={
                "batch_id": batch.id,
                "validation_run_id": val_run.id,
                "run_number": val_run.run_number,
                "total_rows": val_run.total_rows,
                "blocking_errors": val_run.blocking_errors,
                "requiring_review": val_run.requiring_review,
                "declared_discrepancy": val_run.declared_discrepancy
            }
        )
        self.db.commit()

        return {
            "batch_id": batch.id,
            "status": batch.status,
            "validation_run_id": val_run.id,
            "run_number": val_run.run_number,
            "disposition": {
                "total_rows": val_run.total_rows,
                "duplicate_rows": val_run.duplicate_rows,
                "unique_records": val_run.unique_records,
                "blocking_errors": val_run.blocking_errors,
                "requiring_review": val_run.requiring_review,
                "eligible_to_proceed": val_run.eligible_to_proceed,
                "declared_intake": val_run.declared_intake,
                "declared_discrepancy": val_run.declared_discrepancy
            },
            "findings_count": len(findings),
            "findings": [
                {
                    "id": f.id,
                    "validation_run_id": f.validation_run_id,
                    "row_index": f.row_index,
                    "field_name": f.field_name,
                    "finding_type": f.finding_type,
                    "rule_id": f.rule_id,
                    "severity": f.severity,
                    "submitted_value": f.submitted_value,
                    "expected_value": f.expected_value,
                    "reference_source_id": f.reference_source_id,
                    "message": f.message,
                    "recommended_action": f.recommended_action
                }
                for f in findings
            ]
        }

    def get_or_sync_staged_rows(self, batch_id: int) -> List[StagedBatchRow]:
        """
        Retrieves all staged rows for in-browser spreadsheet rendering.
        If no staged rows exist yet, automatically synchronizes them from batch storage.
        """
        batch = self.db.query(UploadBatch).filter(UploadBatch.id == batch_id).first()
        if not batch:
            raise ValueError(f"Batch #{batch_id} not found.")

        staged_rows = self.db.query(StagedBatchRow).filter(
            StagedBatchRow.batch_id == batch_id
        ).order_by(StagedBatchRow.row_index.asc()).all()

        if staged_rows:
            return staged_rows

        mapping_config = batch.mapping_config or {}
        storage_file_path = mapping_config.get("_storage_file")
        if not storage_file_path or not Path(storage_file_path).exists():
            return []

        file_bytes = Path(storage_file_path).read_bytes()
        df = self._load_dataframe(file_bytes, batch.filename)
        canonical_to_source = {v: k for k, v in mapping_config.items() if not k.startswith("_")}

        # Run reconciliation to populate StagedBatchRow
        self.execute_three_way_reconciliation(
            df=df,
            batch=batch,
            user_id=batch.submitted_by_user_id,
            trigger_action="REVALIDATION",
            canonical_to_source=canonical_to_source
        )

        return self.db.query(StagedBatchRow).filter(
            StagedBatchRow.batch_id == batch_id
        ).order_by(StagedBatchRow.row_index.asc()).all()

    def sync_storage_files(self, batch: UploadBatch, df: pd.DataFrame) -> Dict[str, str]:
        """
        Synchronizes both Excel (.xlsx) and CSV (.csv) backing storage files on disk.
        Ensures genuine OpenXML formatting, sets APAAR IDs explicitly as text format,
        and auto-fits columns.
        """
        storage_dir = Path("storage/uploads").resolve()
        storage_dir.mkdir(parents=True, exist_ok=True)

        mapping_config = dict(batch.mapping_config or {})
        base_name = Path(batch.filename).stem or f"batch_{batch.id}"
        file_hash = batch.file_hash or f"sync_{batch.id}"

        xlsx_path_str = mapping_config.get("_storage_file_xlsx")
        csv_path_str = mapping_config.get("_storage_file_csv")

        xlsx_file = Path(xlsx_path_str) if xlsx_path_str else (storage_dir / f"batch_{file_hash[:10]}_{base_name}.xlsx")
        csv_file = Path(csv_path_str) if csv_path_str else (storage_dir / f"batch_{file_hash[:10]}_{base_name}.csv")

        # 1. Write clean CSV
        try:
            df.to_csv(csv_file, index=False, encoding="utf-8")
        except Exception:
            pass

        # 2. Write genuine OpenXML XLSX workbook
        try:
            import openpyxl.styles
            import openpyxl.utils
            with pd.ExcelWriter(xlsx_file, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Admissions_Roster")
                ws = writer.sheets["Admissions_Roster"]
                header_fill = openpyxl.styles.PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
                header_font = openpyxl.styles.Font(color="FFFFFF", bold=True, name="Calibri")
                for col in range(1, len(df.columns) + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

                # Format APAAR ID column as text to prevent scientific notation in Excel
                for col_idx, col_name in enumerate(df.columns, start=1):
                    if "apaar" in str(col_name).lower():
                        for row_idx in range(2, len(df) + 2):
                            ws.cell(row=row_idx, column=col_idx).number_format = '@'
                    # Auto column width
                    col_letter = openpyxl.utils.get_column_letter(col_idx)
                    max_len = max(len(str(val or '')) for val in df[col_name])
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        except Exception:
            try:
                df.to_excel(xlsx_file, index=False)
            except Exception:
                pass

        mapping_config["_storage_file_xlsx"] = str(xlsx_file.as_posix())
        mapping_config["_storage_file_csv"] = str(csv_file.as_posix())
        is_excel = batch.filename.lower().endswith(('.xlsx', '.xls'))
        mapping_config["_storage_file"] = str(xlsx_file.as_posix() if is_excel else csv_file.as_posix())
        mapping_config["_file_format"] = "XLSX" if is_excel else "CSV"
        batch.mapping_config = mapping_config
        self.db.flush()

        return {
            "xlsx": str(xlsx_file.resolve()),
            "csv": str(csv_file.resolve())
        }

    def update_staged_row(
        self,
        batch_id: int,
        row_id: int,
        updates: Dict[str, Any],
        user_id: int,
        user_email: str,
        user_role: str
    ) -> Dict[str, Any]:
        """
        SIH Requirement: In-Browser Spreadsheet Inline Row Correction.
        Updates row in database, rewrites backing storage file, and immediately re-evaluates
        authoritative Python reconciliation algorithms in real time.
        """
        batch = self.db.query(UploadBatch).filter(UploadBatch.id == batch_id).first()
        if not batch:
            raise ValueError(f"Batch #{batch_id} not found.")

        if batch.status in ["APPROVED", "REJECTED"]:
            raise ValueError(f"Cannot edit Batch #{batch_id} because it has already been adjudicated as {batch.status}.")

        staged_row = self.db.query(StagedBatchRow).filter(
            StagedBatchRow.id == row_id,
            StagedBatchRow.batch_id == batch_id
        ).first()
        if not staged_row:
            raise ValueError(f"Staged row #{row_id} not found in Batch #{batch_id}.")

        # Apply cell updates
        if "apaar_id" in updates and updates["apaar_id"] is not None:
            staged_row.apaar_id = str(updates["apaar_id"]).strip()
        if "student_name" in updates and updates["student_name"] is not None:
            staged_row.student_name = str(updates["student_name"]).strip()
        if "programme_code" in updates and updates["programme_code"] is not None:
            staged_row.programme_code = str(updates["programme_code"]).strip().upper()
        if "category" in updates and updates["category"] is not None:
            staged_row.category = str(updates["category"]).strip().upper()
        if "gender" in updates and updates["gender"] is not None:
            staged_row.gender = str(updates["gender"]).strip().upper()
        if "admission_quota" in updates and updates["admission_quota"] is not None:
            staged_row.admission_quota = str(updates["admission_quota"]).strip()
        if "dob" in updates and updates["dob"]:
            try:
                staged_row.dob = pd.to_datetime(updates["dob"]).date()
            except Exception:
                pass

        staged_row.updated_at = datetime.utcnow()
        self.db.flush()

        # Re-assemble complete DataFrame from all staged rows in order
        all_staged = self.db.query(StagedBatchRow).filter(
            StagedBatchRow.batch_id == batch_id
        ).order_by(StagedBatchRow.row_index.asc()).all()

        mapping_config = batch.mapping_config or {}
        canonical_to_source = {v: k for k, v in mapping_config.items() if not k.startswith("_")}
        if not canonical_to_source:
            canonical_to_source = {
                "apaar_id": "apaar_id",
                "student_name": "student_name",
                "programme_code": "programme_code",
                "category": "category",
                "gender": "gender",
                "admission_quota": "admission_quota",
                "dob": "dob"
            }

        rows_data = []
        for r in all_staged:
            row_dict = {}
            row_dict[canonical_to_source.get("apaar_id", "apaar_id")] = r.apaar_id or ""
            row_dict[canonical_to_source.get("student_name", "student_name")] = r.student_name or ""
            row_dict[canonical_to_source.get("programme_code", "programme_code")] = r.programme_code or ""
            row_dict[canonical_to_source.get("category", "category")] = r.category or "OPEN"
            row_dict[canonical_to_source.get("gender", "gender")] = r.gender or "MALE"
            row_dict[canonical_to_source.get("admission_quota", "admission_quota")] = r.admission_quota or "STATE_GOVT"
            row_dict[canonical_to_source.get("dob", "dob")] = r.dob.isoformat() if r.dob else ""
            rows_data.append(row_dict)

        updated_df = pd.DataFrame(rows_data)

        # Synchronize both disk files (CSV and genuine styled Excel XLSX)
        self.sync_storage_files(batch, updated_df)

        # Clear prior findings
        self.db.query(ValidationFinding).filter(ValidationFinding.batch_id == batch_id).delete()
        self.db.flush()

        # Re-run reconciliation
        val_run, findings = self.execute_three_way_reconciliation(
            df=updated_df,
            batch=batch,
            user_id=user_id,
            trigger_action="REVALIDATION",
            canonical_to_source=canonical_to_source
        )

        self.db.refresh(staged_row)
        row_findings = [f for f in findings if f.row_index == staged_row.row_index]

        append_audit_event(
            db=self.db,
            action="STAGED_ROW_INLINE_CORRECTION",
            entity_type="StagedBatchRow",
            entity_id=str(staged_row.id),
            actor_id=user_id,
            actor_email=user_email,
            actor_role=user_role,
            metadata_json={
                "batch_id": batch_id,
                "row_id": staged_row.id,
                "row_index": staged_row.row_index,
                "updates": {k: str(v) for k, v in updates.items() if v is not None},
                "new_status": staged_row.status,
                "findings_count_for_row": len(row_findings)
            }
        )
        self.db.commit()

        return {
            "staged_row": {
                "id": staged_row.id,
                "batch_id": staged_row.batch_id,
                "row_index": staged_row.row_index,
                "apaar_id": staged_row.apaar_id,
                "student_name": staged_row.student_name,
                "programme_code": staged_row.programme_code,
                "category": staged_row.category,
                "gender": staged_row.gender,
                "dob": staged_row.dob.isoformat() if staged_row.dob else None,
                "admission_quota": staged_row.admission_quota,
                "status": staged_row.status,
                "findings": [
                    {
                        "id": f.id,
                        "field_name": f.field_name,
                        "finding_type": f.finding_type,
                        "rule_id": f.rule_id,
                        "severity": f.severity,
                        "message": f.message,
                        "recommended_action": f.recommended_action
                    }
                    for f in row_findings
                ]
            },
            "batch_status": batch.status,
            "validation_run_id": val_run.id,
            "run_number": val_run.run_number,
            "disposition": {
                "total_rows": val_run.total_rows,
                "unique_records": val_run.unique_records,
                "blocking_errors": val_run.blocking_errors,
                "requiring_review": val_run.requiring_review,
                "eligible_to_proceed": val_run.eligible_to_proceed
            }
        }
