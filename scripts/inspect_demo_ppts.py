import pptx
import os
import sys

files = [
    r'D:\AICTE\demoppt\TCE_Yquanto_SIH26263.pptx',
    r'D:\AICTE\demoppt\PRAMAAN-AICTE_SIH2026_Idea_Presentation.pptx',
    r'D:\AICTE\finaldemoppt\SIH2026-IDEA-Presentation-Format.pptx'
]

with open(r'D:\AICTE\demoppt_inspection_dump.txt', 'w', encoding='utf-8') as out:
    for f in files:
        if not os.path.exists(f):
            out.write(f"File not found: {f}\n")
            continue
        prs = pptx.Presentation(f)
        out.write(f"\n=======================================================\n")
        out.write(f"FILE: {f} (Slides: {len(prs.slides)})\n")
        out.write(f"=======================================================\n")
        for idx, slide in enumerate(prs.slides):
            out.write(f"\n--- Slide {idx+1} ---\n")
            for shp in slide.shapes:
                if shp.has_text_frame and shp.text_frame.text.strip():
                    lines = [l.strip() for l in shp.text_frame.text.split('\n') if l.strip()]
                    text_preview = " // ".join(lines[:4])
                    out.write(f"  [{shp.name}] ({shp.left},{shp.top},{shp.width},{shp.height}) -> {text_preview[:150]}\n")
                elif shp.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    out.write(f"  [PICTURE: {shp.name}] ({shp.left},{shp.top},{shp.width},{shp.height})\n")

print("Dumped demo PPTs successfully!")
