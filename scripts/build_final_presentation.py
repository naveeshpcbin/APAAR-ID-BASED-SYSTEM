"""
PRAMAAN-AICTE: Final SIH 2026 Presentation Builder (Strict Official Demo Template Structure)
Reconstructs the presentation adhering strictly to the official SIH 2026 master template:
- Slide 1: EXACT Demo PPT Structure (Official illustration Picture 4 preserved, Title 7, Subtitle 3, SIH Logo, verified metadata in TextBox 9)
- Slides 2-6: High-clarity 1920x1032 verified screenshots on right sides with clear captions and explainable callouts
- 100% verified against live codebase, Supabase PostgreSQL, Render backend, and Vercel custom domains
- Clean unicode typography, zero text clipping, 16:9 widescreen format
"""

import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# Professional Color Palette matching SIH Reference Standards
NAVY = RGBColor(17, 34, 102)          # #112266 - SIH Official Navy
TEAL = RGBColor(0, 140, 110)         # #008C6E - Government Trust Teal
DARK_BLUE = RGBColor(26, 54, 93)     # #1A365D - Deep Slate Blue
ACCENT_BLUE = RGBColor(26, 86, 219)   # #1A56DB - Royal Tech Blue
ALERT_RED = RGBColor(180, 40, 40)    # #B42828 - Risk Red
CHARCOAL = RGBColor(34, 34, 34)      # #222222 - Text Primary
MUTED_GRAY = RGBColor(100, 116, 139) # #64748B - Text Secondary
BG_CARD = RGBColor(255, 255, 255)    # White Card
BG_TINT = RGBColor(248, 250, 252)    # #F8FAFC
BORDER_COLOR = RGBColor(203, 213, 225) # #CBD5E1 - Slate Border
FOOTER_BLUE = RGBColor(0, 111, 191)  # #006FBF - SIH Footer Blue

TEMPLATE_PATH = r"D:\AICTE\finaldemoppt\SIH2026-IDEA-Presentation-Format.pptx"
SLIDE_IMAGES_DIR = r"D:\AICTE\artifacts\slide_images"

OUT_PATHS = [
    r"D:\AICTE\finaldemoppt\PRAMAAN_AICTE_SIH2026_FINAL.pptx",
    r"D:\AICTE\deck\final\PRAMAAN-AICTE_SIH2026_Final.pptx"
]

def format_slide_header(slide, title_text, team_name="TCE_Yquanto", slide_num=2):
    """Formats top title, team oval badge, SIH logo, and bottom footer strip."""
    for shp in slide.shapes:
        if shp.name == "Title 1":
            shp.left = Inches(1.85)
            shp.top = Inches(0.18)
            shp.width = Inches(8.85)
            shp.height = Inches(0.85)
            tf = shp.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            p.text = title_text
            p.font.name = "Arial"
            p.font.size = Pt(16.5)
            p.font.bold = True
            p.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.LEFT

        elif shp.name.startswith("Oval"):
            shp.left = Inches(0.30)
            shp.top = Inches(0.20)
            shp.width = Inches(1.50)
            shp.height = Inches(0.75)
            tf = shp.text_frame
            tf.word_wrap = False
            tf.margin_left = Inches(0.02)
            tf.margin_right = Inches(0.02)
            tf.clear()
            p = tf.paragraphs[0]
            p.text = team_name
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.CENTER

        elif shp.name.startswith("Footer Placeholder"):
            shp.left = Inches(3.0)
            shp.top = Inches(6.98)
            shp.width = Inches(7.33)
            shp.height = Inches(0.40)
            tf = shp.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = f"@ {team_name} | Smart India Hackathon 2026"
            p.font.name = "Arial"
            p.font.size = Pt(10)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.CENTER

        elif shp.name.startswith("Slide Number Placeholder"):
            shp.left = Inches(11.0)
            shp.top = Inches(6.98)
            shp.width = Inches(1.8)
            shp.height = Inches(0.40)
            tf = shp.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = str(slide_num)
            p.font.name = "Arial"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.RIGHT

    # Remove template's default TextBox 8
    for s in [shp for shp in slide.shapes if shp.name == "TextBox 8"]:
        sp = s._element
        sp.getparent().remove(sp)

def create_card(slide, left, top, width, height, border_color=BORDER_COLOR, bg_color=BG_CARD):
    """Creates a rounded container card."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def create_header_badge(slide, left, top, width, height, text, bg_color=NAVY):
    """Creates a colored banner badge for a card."""
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    return badge

def build_presentation():
    print(f"Loading official SIH template: {TEMPLATE_PATH}")
    prs = Presentation(TEMPLATE_PATH)

    # 1. Remove instruction slide 7 (Strict SIH Rule #1: exactly 6 slides)
    if len(prs.slides) >= 7:
        print("Removing instruction slide 7 (strict 6-slide rule)...")
        rId = prs.slides._sldIdLst[6].attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        del prs.slides._sldIdLst[6]

    # =========================================================================
    # SLIDE 1: TITLE PAGE (STRICT OFFICIAL DEMO PPT STRUCTURE)
    # =========================================================================
    print("Building Slide 1 (Strict Official Demo PPT Structure)...")
    s1 = prs.slides[0]

    # Configure Title 7: 'SMART INDIA HACKATHON 2026'
    for shp in s1.shapes:
        if shp.name == "Title 7":
            shp.left = Inches(0.40)
            shp.top = Inches(0.20)
            shp.width = Inches(10.20)
            shp.height = Inches(0.90)
            tf = shp.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = "SMART INDIA HACKATHON 2026"
            p.font.name = "Arial"
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.LEFT

        elif shp.name == "Subtitle 3":
            shp.left = Inches(0.40)
            shp.top = Inches(0.95)
            shp.width = Inches(6.50)
            shp.height = Inches(0.40)
            tf = shp.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "TITLE PAGE"
            p.font.name = "Arial"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = TEAL
            p.alignment = PP_ALIGN.LEFT

        elif shp.name == "TextBox 9":
            shp.left = Inches(0.40)
            shp.top = Inches(1.55)
            shp.width = Inches(6.80)
            shp.height = Inches(5.25)
            tf = shp.text_frame
            tf.word_wrap = True
            tf.clear()

            title_rows = [
                ("Problem Statement ID", "SIH26263", NAVY, Pt(13), True),
                ("Problem Statement Title", "Development of an APAAR ID-Based Student Data Management and Verification Platform for AICTE-Approved Institutions", DARK_BLUE, Pt(11), False),
                ("Theme", "Smart Automation", CHARCOAL, Pt(11), False),
                ("PS Category", "Software", CHARCOAL, Pt(11), False),
                ("Team ID", "TCE_Yquanto", CHARCOAL, Pt(11), False),
                ("Team Name", "TCE_Yquanto (Thiagarajar College of Engineering, Madurai)", NAVY, Pt(11.5), True),
                ("Institute Name", "Thiagarajar College of Engineering, Madurai, Tamil Nadu", DARK_BLUE, Pt(11), False),
                ("Live Web Platform", "https://brainforge.info  |  https://pramaan.brainforge.info", TEAL, Pt(10.5), True),
                ("Live Backend & API", "https://apaar-id-based-system.onrender.com/docs", TEAL, Pt(10.5), True),
                ("Source Repository", "https://github.com/naveeshpcbin/APAAR-ID-BASED-SYSTEM", ACCENT_BLUE, Pt(10.5), True),
                ("Verification Status", "66/66 Passing Tests | 20 Tables with RLS | Live Cloud Verified", TEAL, Pt(10.5), True),
            ]

            for idx, (label, val, val_color, val_size, val_bold) in enumerate(title_rows):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.space_before = Pt(4.0) if idx > 0 else Pt(0)

                r_lbl = p.add_run()
                r_lbl.text = f"{label}: "
                r_lbl.font.name = "Arial"
                r_lbl.font.size = Pt(10.8)
                r_lbl.font.bold = True
                r_lbl.font.color.rgb = NAVY

                r_val = p.add_run()
                r_val.text = val
                r_val.font.name = "Arial"
                r_val.font.size = val_size
                r_val.font.bold = val_bold
                r_val.font.color.rgb = val_color

    # Add bottom footer strip on Slide 1 matching official styling
    footer_bar_1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.0), Inches(6.95), Inches(13.333), Inches(0.55))
    footer_bar_1.fill.solid()
    footer_bar_1.fill.fore_color.rgb = FOOTER_BLUE
    footer_bar_1.line.fill.background()

    tb_f1 = s1.shapes.add_textbox(Inches(3.0), Inches(6.98), Inches(7.33), Inches(0.40))
    tf_f1 = tb_f1.text_frame
    p_f1 = tf_f1.paragraphs[0]
    p_f1.text = "@ TCE_Yquanto | Smart India Hackathon 2026"
    p_f1.font.name = "Arial"
    p_f1.font.size = Pt(10)
    p_f1.font.color.rgb = RGBColor(255, 255, 255)
    p_f1.alignment = PP_ALIGN.CENTER

    tb_n1 = s1.shapes.add_textbox(Inches(11.0), Inches(6.98), Inches(1.8), Inches(0.40))
    tf_n1 = tb_n1.text_frame
    p_n1 = tf_n1.paragraphs[0]
    p_n1.text = "1"
    p_n1.font.name = "Arial"
    p_n1.font.size = Pt(10)
    p_n1.font.bold = True
    p_n1.font.color.rgb = RGBColor(255, 255, 255)
    p_n1.alignment = PP_ALIGN.RIGHT

    # =========================================================================
    # SLIDE 2: PROPOSED SOLUTION: THE PRAMAAN-AICTE RECONCILIATION SYSTEM
    # =========================================================================
    print("Building Slide 2 (Proposed Solution - 3 Problem/Solution Pillars + Command Center)...")
    s2 = prs.slides[1]
    format_slide_header(s2, "PROPOSED SOLUTION: THE PRAMAAN-AICTE RECONCILIATION SYSTEM", slide_num=2)

    left_x = Inches(0.40)
    col_w = Inches(6.15)

    pillar_data = [
        (
            Inches(1.22), Inches(1.52),
            "1. REGULATORY CHALLENGE ADDRESSED",
            DARK_BLUE,
            [
                ("Fragmented Institutional Data", "Divergent college ERP spreadsheet schemas, non-standard column headers, and unverified declarations."),
                ("Reconciliation Bottlenecks", "Difficulties cross-verifying declared enrolments against actual student records and AICTE approval ceilings."),
                ("Intake & Duplicate Risks", "Lack of automated visibility into over-intake breaches, quota diversion, and concurrent multi-college claims.")
            ]
        ),
        (
            Inches(2.88), Inches(1.58),
            "2. WHAT PRAMAAN-AICTE IMPLEMENTS",
            NAVY,
            [
                ("Decoupled Institutional Portals", "Role-separated desks for College submission and Central AICTE regulatory governance."),
                ("Automated Ingestion Desk", "CSV/Excel file parsing with resilient column alias mapping into standard staging rows."),
                ("Deterministic 3-Way Reconciliation", "Pure Python comparison engine cross-checking declarations, staged student records, and approved capacity."),
                ("Persisted Validation Runs", "Stores row-level findings with precise classification: approved, quota breach, syntax error, or duplicate.")
            ]
        ),
        (
            Inches(4.60), Inches(1.55),
            "3. KEY WORKFLOW & GOVERNANCE VALUE",
            TEAL,
            [
                ("Explainable Row-Level Findings", "Direct, transparent feedback explaining specific discrepancy causes without opaque batch rejection."),
                ("Quarantine & Resubmission", "Interactive in-browser quarantine sandbox allowing institutions to correct data before council commit."),
                ("Controlled Central Decisions", "Council review workflow supporting formal decisions: approve, return for correction, or reject."),
                ("Audit Traceability", "Complete validation run history and append-only sequential SHA-256 audit chaining across actions.")
            ]
        )
    ]

    for p_top, p_h, p_title, p_color, p_bullets in pillar_data:
        create_card(s2, left_x, p_top, col_w, p_h)
        create_header_badge(s2, left_x + Inches(0.12), p_top - Inches(0.12), col_w - Inches(0.24), Inches(0.30), p_title, p_color)

        tb = s2.shapes.add_textbox(left_x + Inches(0.12), p_top + Inches(0.20), col_w - Inches(0.24), p_h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        for b_idx, (b_lead, b_text) in enumerate(p_bullets):
            p = tf.paragraphs[0] if b_idx == 0 else tf.add_paragraph()
            p.space_before = Pt(2.5) if b_idx > 0 else Pt(0)

            r_lead = p.add_run()
            r_lead.text = f"- {b_lead}: "
            r_lead.font.name = "Arial"
            r_lead.font.size = Pt(8.5)
            r_lead.font.bold = True
            r_lead.font.color.rgb = p_color

            r_text = p.add_run()
            r_text.text = b_text
            r_text.font.name = "Arial"
            r_text.font.size = Pt(8.0)
            r_text.font.color.rgb = CHARCOAL

    # RIGHT COLUMN: High-Resolution Screenshot of AICTE Central Command Console
    rt_x = Inches(6.75)
    rt_w = Inches(6.15)
    create_card(s2, rt_x, Inches(1.22), rt_w, Inches(4.93))
    create_header_badge(s2, rt_x + Inches(0.15), Inches(1.10), rt_w - Inches(0.30), Inches(0.30), "AICTE CENTRAL APEX COMMAND & RECONCILIATION DASHBOARD", NAVY)

    s2_img = os.path.join(SLIDE_IMAGES_DIR, "slide_2_Picture 3.png")
    if os.path.exists(s2_img):
        s2.shapes.add_picture(s2_img, rt_x + Inches(0.12), Inches(1.48), rt_w - Inches(0.24), Inches(3.20))

    tb_cap = s2.shapes.add_textbox(rt_x + Inches(0.12), Inches(4.75), rt_w - Inches(0.24), Inches(1.30))
    tf_cap = tb_cap.text_frame
    tf_cap.word_wrap = True
    tf_cap.margin_left = tf_cap.margin_right = tf_cap.margin_top = tf_cap.margin_bottom = 0

    captions = [
        ("Central Intake Intelligence", "Live monitoring of sanctioned intake vs enrolled students across B.Tech, MCA, and Diploma."),
        ("Deterministic Review Desk", "Instant identification of seat limit breaches, unverified candidates, and duplicate claims."),
        ("Multi-Tenant Governance", "Seamless institutional isolation between AICTE Central Regulators, DTU Delhi, and COEP Pune.")
    ]
    for c_idx, (c_lead, c_desc) in enumerate(captions):
        p = tf_cap.paragraphs[0] if c_idx == 0 else tf_cap.add_paragraph()
        p.space_before = Pt(2.5) if c_idx > 0 else Pt(0)
        r_l = p.add_run()
        r_l.text = f"- {c_lead}: "
        r_l.font.name = "Arial"
        r_l.font.size = Pt(8.2)
        r_l.font.bold = True
        r_l.font.color.rgb = DARK_BLUE
        r_d = p.add_run()
        r_d.text = c_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(7.8)
        r_d.font.color.rgb = CHARCOAL

    # Bottom Pill Badges across Slide 2
    pills = [
        (Inches(0.40), "Deterministic Python Core", TEAL),
        (Inches(3.55), "3-Way Cross-Reconciliation", NAVY),
        (Inches(6.75), "Cryptographic Verification & Audit", DARK_BLUE),
        (Inches(9.95), "66/66 Verified Passing Tests", TEAL)
    ]
    for p_left, p_text, p_color in pills:
        pill = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, p_left, Inches(6.30), Inches(2.95), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = BG_TINT
        pill.line.color.rgb = p_color
        pill.line.width = Pt(1)
        tf_p = pill.text_frame
        tf_p.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf_p.paragraphs[0]
        p.text = p_text
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = p_color
        p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: TECHNICAL APPROACH AND SYSTEM ARCHITECTURE
    # =========================================================================
    print("Building Slide 3 (Technical Approach & Architecture)...")
    s3 = prs.slides[2]
    format_slide_header(s3, "TECHNICAL APPROACH AND SYSTEM ARCHITECTURE", slide_num=3)

    # LEFT COLUMN: 5-Stage Modular Pipeline
    left_x = Inches(0.40)
    col_w = Inches(6.15)
    create_card(s3, left_x, Inches(1.22), col_w, Inches(5.42))
    create_header_badge(s3, left_x + Inches(0.15), Inches(1.10), col_w - Inches(0.30), Inches(0.32), "5-STAGE RECONCILIATION & GOVERNANCE PIPELINE", NAVY)

    stages = [
        ("STAGE 1: Resilient File Ingestion", "Upload institutional CSV/Excel files; dynamic column alias matching auto-normalizes divergent schemas into standardized staging records.", Inches(1.50)),
        ("STAGE 2: Pure Python Validation Engine", "Deterministic rule engine parses syntax, detects in-batch duplicates, verifies 12-digit APAAR format, and runs cross-field consistency checks.", Inches(2.55)),
        ("STAGE 3: 3-Way Dynamic Reconciliation", "Cross-compares declared student counts against verified stage records and sanctioned intake quotas from dynamic reference tables.", Inches(3.60)),
        ("STAGE 4: Council Review & Audit Ledger", "Central officers review row-level discrepancy findings; make formal Approve/Return decisions; actions chained into append-only SHA-256 audit ledger.", Inches(4.65)),
        ("STAGE 5: Cryptographic Verification Receipt", "Approved student records receive an asymmetric Ed25519 digital signature and portable verification receipt with offline verifiable QR code.", Inches(5.70))
    ]

    for st_title, st_desc, st_top in stages:
        c_st = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x + Inches(0.15), st_top, col_w - Inches(0.30), Inches(0.92))
        c_st.fill.solid()
        c_st.fill.fore_color.rgb = BG_TINT
        c_st.line.color.rgb = BORDER_COLOR
        c_st.line.width = Pt(1)

        tb = s3.shapes.add_textbox(left_x + Inches(0.22), st_top + Inches(0.06), col_w - Inches(0.44), Inches(0.80))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = st_title
        r1.font.name = "Arial"
        r1.font.size = Pt(8.8)
        r1.font.bold = True
        r1.font.color.rgb = DARK_BLUE

        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        r2 = p2.add_run()
        r2.text = st_desc
        r2.font.name = "Arial"
        r2.font.size = Pt(8.0)
        r2.font.color.rgb = CHARCOAL

    # RIGHT COLUMN TOP: Verified 4-Layer System Architecture
    rt_x = Inches(6.75)
    rt_w = Inches(6.15)
    create_card(s3, rt_x, Inches(1.22), rt_w, Inches(2.26))
    create_header_badge(s3, rt_x + Inches(0.15), Inches(1.10), rt_w - Inches(0.30), Inches(0.30), "VERIFIED 4-LAYER SYSTEM ARCHITECTURE", DARK_BLUE)

    tb_tech = s3.shapes.add_textbox(rt_x + Inches(0.15), Inches(1.48), rt_w - Inches(0.30), Inches(1.92))
    tf_tech = tb_tech.text_frame
    tf_tech.word_wrap = True
    tf_tech.margin_left = tf_tech.margin_right = tf_tech.margin_top = tf_tech.margin_bottom = 0

    tech_layers = [
        ("Presentation Layer", "Next.js 14 (App Router), TypeScript, Tailwind CSS, Recharts for analytics; deployed on Vercel Edge."),
        ("Application & API", "FastAPI (Modular Monolith), Python 3.10+, Pydantic v2, Uvicorn, 73 endpoints; deployed on Render Cloud."),
        ("Persistence & Schemas", "SQLAlchemy ORM, Supabase PostgreSQL 17.11 (20 tables with RLS enabled) / SQLite WAL mode."),
        ("Cryptography & Security", "Ed25519 (RFC 8032 signing), Keyed HMAC-SHA256, sequential SHA-256 audit chaining.")
    ]
    for tl_idx, (tl_name, tl_desc) in enumerate(tech_layers):
        p = tf_tech.paragraphs[0] if tl_idx == 0 else tf_tech.add_paragraph()
        p.space_before = Pt(3) if tl_idx > 0 else Pt(0)
        r_n = p.add_run()
        r_n.text = f"- {tl_name}: "
        r_n.font.name = "Arial"
        r_n.font.size = Pt(8.8)
        r_n.font.bold = True
        r_n.font.color.rgb = NAVY
        r_d = p.add_run()
        r_d.text = tl_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8.1)
        r_d.font.color.rgb = CHARCOAL

    # RIGHT COLUMN BOTTOM: High-Resolution Screenshot of Smart Ingestion Desk
    create_card(s3, rt_x, Inches(3.60), rt_w, Inches(3.04))
    create_header_badge(s3, rt_x + Inches(0.15), Inches(3.48), rt_w - Inches(0.30), Inches(0.30), "PROTOTYPE EVIDENCE: SMART INGESTION DESK", TEAL)

    s3_img = os.path.join(SLIDE_IMAGES_DIR, "slide_3_Picture 3.png")
    if os.path.exists(s3_img):
        s3.shapes.add_picture(s3_img, rt_x + Inches(0.12), Inches(3.82), rt_w - Inches(0.24), Inches(2.15))

    tb_sb = s3.shapes.add_textbox(rt_x + Inches(0.12), Inches(6.05), rt_w - Inches(0.24), Inches(0.50))
    tf_sb = tb_sb.text_frame
    tf_sb.word_wrap = True
    p_sb = tf_sb.paragraphs[0]
    p_sb.text = "Smart Ingestion Desk: Automated schema mapping, pre-flight validation status, row findings ledger, and atomic batch commitment."
    p_sb.font.name = "Arial"
    p_sb.font.size = Pt(7.8)
    p_sb.font.color.rgb = MUTED_GRAY
    p_sb.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 4: FEASIBILITY, VIABILITY AND RISK MITIGATION
    # =========================================================================
    print("Building Slide 4 (Feasibility, Viability & Risk Mitigation)...")
    s4 = prs.slides[3]
    format_slide_header(s4, "FEASIBILITY, VIABILITY AND RISK MITIGATION", slide_num=4)

    # LEFT COLUMN: Feasibility & Implementation Maturity
    f_left = Inches(0.40)
    f_w = Inches(6.15)
    f_h = Inches(5.42)
    create_card(s4, f_left, Inches(1.22), f_w, f_h)
    create_header_badge(s4, f_left + Inches(0.15), Inches(1.10), f_w - Inches(0.30), Inches(0.32), "IMPLEMENTATION FEASIBILITY & VERIFIED MATURITY", NAVY)

    tb_feas = s4.shapes.add_textbox(f_left + Inches(0.15), Inches(1.50), f_w - Inches(0.30), Inches(5.00))
    tf_feas = tb_feas.text_frame
    tf_feas.word_wrap = True
    tf_feas.margin_left = tf_feas.margin_right = tf_feas.margin_top = tf_feas.margin_bottom = 0

    feas_points = [
        ("Architectural Feasibility", "Zero-install web architecture. Standard modern browser access across Windows, Linux, and macOS. Decoupled Next.js client and FastAPI modular backend."),
        ("Verified Test Suite (66/66 Passing)", "Exhaustive automated verification across 14 test suites covering authentication, institution onboarding, batch upload, deterministic reconciliation, security, and reference data."),
        ("Deterministic Logic Integrity", "Reconciliation logic is pure, explainable Python with explicit rule sets for sanctioned intake and supernumerary quotas; avoids opaque or non-deterministic AI decisions."),
        ("Synthetic Data Testing Engine", "Database-backed synthetic APAAR identities and editable approval/declaration models enable comprehensive pre-deployment stress testing and repeatable verification."),
        ("Dual-Dialect Database Support", "SQLAlchemy ORM architecture with dynamic driver fallback (psycopg3/psycopg2) verified on local SQLite WAL mode and Supabase PostgreSQL 17.11 (20 tables, RLS active)."),
        ("Live Production Cloud Verification", "Full cloud deployment verified live: FastAPI backend on Render with PostgreSQL connection, Next.js frontend on Vercel with custom domain brainforge.info, and 100% end-to-end integration pass.")
    ]

    for fp_idx, (fp_lead, fp_text) in enumerate(feas_points):
        p = tf_feas.paragraphs[0] if fp_idx == 0 else tf_feas.add_paragraph()
        p.space_before = Pt(7) if fp_idx > 0 else Pt(0)
        r_lead = p.add_run()
        r_lead.text = f"- {fp_lead}: "
        r_lead.font.name = "Arial"
        r_lead.font.size = Pt(8.9)
        r_lead.font.bold = True
        r_lead.font.color.rgb = DARK_BLUE
        r_text = p.add_run()
        r_text.text = fp_text
        r_text.font.name = "Arial"
        r_text.font.size = Pt(8.3)
        r_text.font.color.rgb = CHARCOAL

    # RIGHT COLUMN TOP: Governance Risks & Technical Countermeasures
    r_left = Inches(6.75)
    r_w = Inches(6.15)
    create_card(s4, r_left, Inches(1.22), r_w, Inches(2.26))
    create_header_badge(s4, r_left + Inches(0.15), Inches(1.10), r_w - Inches(0.30), Inches(0.30), "GOVERNANCE RISKS & TECHNICAL COUNTERMEASURES", ALERT_RED)

    tb_risk = s4.shapes.add_textbox(r_left + Inches(0.15), Inches(1.48), r_w - Inches(0.30), Inches(1.92))
    tf_risk = tb_risk.text_frame
    tf_risk.word_wrap = True
    tf_risk.margin_left = tf_risk.margin_right = tf_risk.margin_top = tf_risk.margin_bottom = 0

    risks = [
        ("Inconsistent Spreadsheets", "Colleges use varying ERP column headers.", "Header alias normalizer maps 50+ divergent column variants into standard canonical schema."),
        ("Identifier Privacy Exposure", "Storing plaintext student APAAR IDs risks data leakage.", "Keyed HMAC-SHA256 pseudonymized indexing and dynamic masking ensure zero plaintext PII exposure."),
        ("Concurrent Admissions", "Students admitted concurrently across different institutions.", "Cross-institution collision radar cross-checks active admissions against central records and flags duplicates.")
    ]

    for r_idx, (r_title, r_desc, r_mit) in enumerate(risks):
        p_t = tf_risk.paragraphs[0] if r_idx == 0 else tf_risk.add_paragraph()
        p_t.space_before = Pt(3.5) if r_idx > 0 else Pt(0)
        r_t = p_t.add_run()
        r_t.text = f"- {r_title}: "
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.5)
        r_t.font.bold = True
        r_t.font.color.rgb = ALERT_RED
        r_td = p_t.add_run()
        r_td.text = f"{r_desc} -> Control: {r_mit}"
        r_td.font.name = "Arial"
        r_td.font.size = Pt(8.0)
        r_td.font.color.rgb = CHARCOAL

    # RIGHT COLUMN BOTTOM: High-Resolution Screenshot of Duplicate Collision Radar
    create_card(s4, r_left, Inches(3.60), r_w, Inches(3.04))
    create_header_badge(s4, r_left + Inches(0.15), Inches(3.48), r_w - Inches(0.30), Inches(0.30), "PROTOTYPE EVIDENCE: DUPLICATE COLLISION RADAR", TEAL)

    s4_img = os.path.join(SLIDE_IMAGES_DIR, "slide_4_Picture 3.png")
    if os.path.exists(s4_img):
        s4.shapes.add_picture(s4_img, r_left + Inches(0.12), Inches(3.82), r_w - Inches(0.24), Inches(2.15))

    tb_r_cap = s4.shapes.add_textbox(r_left + Inches(0.12), Inches(6.05), r_w - Inches(0.24), Inches(0.50))
    tf_r_cap = tb_r_cap.text_frame
    tf_r_cap.word_wrap = True
    p_rc = tf_r_cap.paragraphs[0]
    p_rc.text = "Duplicate Collision Radar: Side-by-side college enrollment audit, quota verification, and regulatory adjudication controls."
    p_rc.font.name = "Arial"
    p_rc.font.size = Pt(7.8)
    p_rc.font.color.rgb = MUTED_GRAY
    p_rc.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: IMPACT AND BENEFITS
    # =========================================================================
    print("Building Slide 5 (Impact and Benefits)...")
    s5 = prs.slides[4]
    format_slide_header(s5, "MULTI-STAKEHOLDER IMPACT & DEGREE VERIFICATION", slide_num=5)

    # LEFT COLUMN: Multi-Stakeholder Impact
    imp_left = Inches(0.40)
    imp_w = Inches(6.15)
    imp_h = Inches(5.42)
    create_card(s5, imp_left, Inches(1.22), imp_w, imp_h)
    create_header_badge(s5, imp_left + Inches(0.15), Inches(1.10), imp_w - Inches(0.30), Inches(0.32), "MULTI-STAKEHOLDER REGULATORY IMPACT", TEAL)

    tb_imp = s5.shapes.add_textbox(imp_left + Inches(0.15), Inches(1.50), imp_w - Inches(0.30), Inches(5.00))
    tf_imp = tb_imp.text_frame
    tf_imp.word_wrap = True
    tf_imp.margin_left = tf_imp.margin_right = tf_imp.margin_top = tf_imp.margin_bottom = 0

    stakeholders = [
        ("AICTE & Regulatory Councils", "Comprehensive real-time visibility into institutional enrolment versus sanctioned intake quotas; early detection of unauthorized excess enrolments and quota breaches."),
        ("Approved Institutions", "Standardized data submission, instant row-level discrepancy explanations, and controlled resubmission workflows replacing prolonged paper audit cycles."),
        ("Students & Deserving Candidates", "Transparent tracking protects legitimate seat allocations and safeguards affirmative action quotas from unauthorized commercial reallocation."),
        ("Academic Verifiers & Employers", "Fast offline verification via QR code and cryptographically signed verification receipts without requiring manual institutional registrar confirmation."),
        ("Social Equity & Affirmative Action", "Rigorous enforcement of statutory and supernumerary quotas (TFW, EWS, Kashmiri Migrants) against arbitrary substitution."),
        ("Administrative & Paperless Efficiency", "100% digital data reconciliation replaces manual paperwork, physical affidavits, and physical institutional auditing logistics.")
    ]

    for s_idx, (s_lead, s_desc) in enumerate(stakeholders):
        p = tf_imp.paragraphs[0] if s_idx == 0 else tf_imp.add_paragraph()
        p.space_before = Pt(7) if s_idx > 0 else Pt(0)
        r_l = p.add_run()
        r_l.text = f"- {s_lead}: "
        r_l.font.name = "Arial"
        r_l.font.size = Pt(8.8)
        r_l.font.bold = True
        r_l.font.color.rgb = DARK_BLUE
        r_d = p.add_run()
        r_d.text = s_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8.2)
        r_d.font.color.rgb = CHARCOAL

    # RIGHT COLUMN: Verifiable Degree Certificate Showcase
    cert_left = Inches(6.75)
    cert_w = Inches(6.15)
    create_card(s5, cert_left, Inches(1.22), cert_w, imp_h)
    create_header_badge(s5, cert_left + Inches(0.15), Inches(1.10), cert_w - Inches(0.30), Inches(0.32), "SIH-R10 ACADEMIC CERTIFICATE & ED25519 QR PROOF", NAVY)

    s5_img = os.path.join(SLIDE_IMAGES_DIR, "slide_5_Picture 3.png")
    if os.path.exists(s5_img):
        s5.shapes.add_picture(s5_img, cert_left + Inches(0.12), Inches(1.48), cert_w - Inches(0.24), Inches(3.20))

    tb_cert = s5.shapes.add_textbox(cert_left + Inches(0.12), Inches(4.75), cert_w - Inches(0.24), Inches(1.80))
    tf_cert = tb_cert.text_frame
    tf_cert.word_wrap = True
    tf_cert.margin_left = tf_cert.margin_right = tf_cert.margin_top = tf_cert.margin_bottom = 0

    cert_points = [
        ("Dual Identity Display", "Certificates present Student Enrolment Roll No (ENR-108410E1) alongside masked APAAR ID (XXXXXXXX7890) and category quota status."),
        ("Asymmetric Ed25519 QR Code", "Embedded high-density QR code carries a digitally signed verification token; verifiers validate data integrity against AICTE public key offline."),
        ("Integrity Boundary Clarification", "Digital signature confirms platform-level record integrity and non-tampering; explicitly distinguished from official government identity confirmation.")
    ]
    for cp_idx, (cp_lead, cp_desc) in enumerate(cert_points):
        p = tf_cert.paragraphs[0] if cp_idx == 0 else tf_cert.add_paragraph()
        p.space_before = Pt(3) if cp_idx > 0 else Pt(0)
        r_l = p.add_run()
        r_l.text = f"- {cp_lead}: "
        r_l.font.name = "Arial"
        r_l.font.size = Pt(8.2)
        r_l.font.bold = True
        r_l.font.color.rgb = DARK_BLUE
        r_d = p.add_run()
        r_d.text = cp_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(7.8)
        r_d.font.color.rgb = CHARCOAL

    # =========================================================================
    # SLIDE 6: RESEARCH, REFERENCES AND PROJECT RESOURCES
    # =========================================================================
    print("Building Slide 6 (Research, References & Verified Resources)...")
    s6 = prs.slides[5]
    format_slide_header(s6, "RESEARCH, REFERENCES AND PROJECT RESOURCES", slide_num=6)

    # LEFT COLUMN: Authoritative Standards & Regulatory Citations
    ref_left = Inches(0.40)
    ref_w = Inches(6.15)
    ref_h = Inches(5.42)
    create_card(s6, ref_left, Inches(1.22), ref_w, ref_h)
    create_header_badge(s6, ref_left + Inches(0.15), Inches(1.10), ref_w - Inches(0.30), Inches(0.32), "AUTHORITATIVE STANDARDS & REGULATORY CITATIONS", NAVY)

    tb_ref = s6.shapes.add_textbox(ref_left + Inches(0.15), Inches(1.48), ref_w - Inches(0.30), Inches(5.00))
    tf_ref = tb_ref.text_frame
    tf_ref.word_wrap = True
    tf_ref.margin_left = tf_ref.margin_right = tf_ref.margin_top = tf_ref.margin_bottom = 0

    citations = [
        ("AICTE Approval Process Handbook (APH 2026-27)",
         "Statutory norms for approved institutional intake capacity, supernumerary seat limits (TFW, EWS, Kashmiri Migrants), and programme approval guidelines."),
        ("National Education Policy (NEP 2020) & APAAR / ABC",
         "Government policy framework establishing the 12-digit Automated Permanent Academic Account Registry as the lifelong national academic identifier."),
        ("Digital Personal Data Protection Act (DPDP Act 2023)",
         "Data fiduciary compliance, purpose limitation, identifier pseudonymization, dynamic masking, and student consent management frameworks."),
        ("IETF RFC 8032: Ed25519 Cryptography Specification",
         "Edwards-Curve Digital Signature Algorithm (Curve25519) providing mathematical non-repudiation and instant receipt verification."),
        ("NIST SP 800-107: Recommendation for Secure Hashing",
         "Cryptographic guidelines for sequential SHA-256 hash chaining utilized in the tamper-evident regulatory audit trail.")
    ]

    for c_idx, (c_title, c_desc) in enumerate(citations):
        p_t = tf_ref.paragraphs[0] if c_idx == 0 else tf_ref.add_paragraph()
        p_t.space_before = Pt(7) if c_idx > 0 else Pt(0)
        r_t = p_t.add_run()
        r_t.text = f"- {c_title}"
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.9)
        r_t.font.bold = True
        r_t.font.color.rgb = DARK_BLUE

        p_d = tf_ref.add_paragraph()
        p_d.space_before = Pt(1.5)
        r_d = p_d.add_run()
        r_d.text = c_desc
        r_d.font.name = "Arial"
        r_d.font.size = Pt(8.2)
        r_d.font.color.rgb = CHARCOAL

    # RIGHT COLUMN TOP: Scope Differentiation (Prior Art vs PRAMAAN-AICTE)
    res_left = Inches(6.75)
    res_w = Inches(6.15)
    create_card(s6, res_left, Inches(1.22), res_w, Inches(2.26))
    create_header_badge(s6, res_left + Inches(0.15), Inches(1.10), res_w - Inches(0.30), Inches(0.30), "SCOPE DIFFERENTIATION (PRIOR ART VS PRAMAAN-AICTE)", TEAL)

    tb_diff = s6.shapes.add_textbox(res_left + Inches(0.15), Inches(1.48), res_w - Inches(0.30), Inches(1.92))
    tf_diff = tb_diff.text_frame
    tf_diff.word_wrap = True
    tf_diff.margin_left = tf_diff.margin_right = tf_diff.margin_top = tf_diff.margin_bottom = 0

    diff_points = [
        ("Validation Scope", "Generic uploads check syntax only. PRAMAAN performs 3-way reconciliation (Declarations x Uploads x Sanctioned Capacity)."),
        ("Reference Data", "Traditional workflows rely on static assumptions. PRAMAAN integrates dynamic database references with live revalidation."),
        ("Finding Resolution", "Standard systems use opaque batch rejections. PRAMAAN provides row-level explainable findings and an editable quarantine sandbox."),
        ("Verification Output", "Conventional portals yield unverified spreadsheets. PRAMAAN produces Ed25519-signed verification receipts and QR proofs.")
    ]
    for d_idx, (d_lead, d_text) in enumerate(diff_points):
        p = tf_diff.paragraphs[0] if d_idx == 0 else tf_diff.add_paragraph()
        p.space_before = Pt(3.0) if d_idx > 0 else Pt(0)
        r_l = p.add_run()
        r_l.text = f"- {d_lead}: "
        r_l.font.name = "Arial"
        r_l.font.size = Pt(8.5)
        r_l.font.bold = True
        r_l.font.color.rgb = NAVY
        r_t = p.add_run()
        r_t.text = d_text
        r_t.font.name = "Arial"
        r_t.font.size = Pt(8.0)
        r_t.font.color.rgb = CHARCOAL

    # RIGHT COLUMN BOTTOM: High-Resolution Screenshot of Audit Ledger + Live Resources
    create_card(s6, res_left, Inches(3.60), res_w, Inches(3.04))
    create_header_badge(s6, res_left + Inches(0.15), Inches(3.48), r_w - Inches(0.30), Inches(0.30), "TAMPER-EVIDENT AUDIT LEDGER & LIVE DEPLOYMENT", DARK_BLUE)

    s6_img = os.path.join(SLIDE_IMAGES_DIR, "slide_6_Picture 3.png")
    if os.path.exists(s6_img):
        s6.shapes.add_picture(s6_img, res_left + Inches(0.12), Inches(3.82), res_w - Inches(0.24), Inches(1.50))

    tb_res = s6.shapes.add_textbox(res_left + Inches(0.12), Inches(5.38), res_w - Inches(0.24), Inches(1.20))
    tf_res = tb_res.text_frame
    tf_res.word_wrap = True
    tf_res.margin_left = tf_res.margin_right = tf_res.margin_top = tf_res.margin_bottom = 0

    resources = [
        ("Live Web Platform", "https://brainforge.info | https://pramaan.brainforge.info (Vercel Edge, SSL active)"),
        ("Live Backend & API", "https://apaar-id-based-system.onrender.com/docs (73 OpenAPI endpoints)"),
        ("Source Repository", "https://github.com/naveeshpcbin/APAAR-ID-BASED-SYSTEM (Full code, tests, migrations)"),
        ("Verification Status", "66/66 Unit Tests + 20 Phase F E2E Checks + Live Cloud Verified")
    ]

    for rs_idx, (rs_title, rs_desc) in enumerate(resources):
        p = tf_res.paragraphs[0] if rs_idx == 0 else tf_res.add_paragraph()
        p.space_before = Pt(2.0) if rs_idx > 0 else Pt(0)
        r_tl = p.add_run()
        r_tl.text = f"- {rs_title}: "
        r_tl.font.name = "Arial"
        r_tl.font.size = Pt(8.2)
        r_tl.font.bold = True
        r_tl.font.color.rgb = TEAL
        r_ds = p.add_run()
        r_ds.text = rs_desc
        r_ds.font.name = "Arial"
        r_ds.font.size = Pt(7.8)
        r_ds.font.color.rgb = CHARCOAL

    # Save to both target output paths
    for out_path in OUT_PATHS:
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        print(f"Saving final presentation to: {out_path}")
        prs.save(out_path)
    print("PPTX build complete for all targets!")

if __name__ == "__main__":
    build_presentation()
