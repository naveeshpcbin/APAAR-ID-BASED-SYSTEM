import pptx
from pptx import Presentation
import os

files = [
    r"D:\AICTE\finaldemoppt\SIH2026-IDEA-Presentation-Format.pptx",
    r"D:\AICTE\finaldemoppt\PRAMAAN-AICTE_SIH2026_Idea_Presentation.pptx",
    r"D:\AICTE\demoppt\TCE_Yquanto_SIH26263.pptx"
]

for fpath in files:
    if not os.path.exists(fpath):
        continue
    prs = Presentation(fpath)
    print(f"==================================================")
    print(f"FILE: {os.path.basename(fpath)}")
    print(f"Slide width={prs.slide_width.inches:.3f} in, height={prs.slide_height.inches:.3f} in")
    print(f"Total slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        print(f"  Slide {i+1}: shapes={len(slide.shapes)}")
        shape_types = {}
        for s in slide.shapes:
            shape_types[s.shape_type] = shape_types.get(s.shape_type, 0) + 1
        print(f"     shape types: {shape_types}")
        # print titles or top text
        titles = []
        for s in slide.shapes:
            if s.has_text_frame and s.text_frame.text.strip():
                t = s.text_frame.text.strip().replace('\n', ' ')
                if len(t) < 80:
                    titles.append(t)
        print(f"     text previews: {titles[:3]}")

