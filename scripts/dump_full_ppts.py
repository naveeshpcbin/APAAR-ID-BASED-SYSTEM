import pptx

def dump_full(pptx_path, out_file):
    prs = pptx.Presentation(pptx_path)
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(f"FULL DUMP OF {pptx_path}\n")
        f.write(f"Slide Width: {prs.slide_width}, Height: {prs.slide_height}\n\n")
        for i, slide in enumerate(prs.slides):
            f.write(f"==================== SLIDE {i+1} ====================\n")
            for shp in slide.shapes:
                f.write(f"Shape: {shp.name} (type: {shp.shape_type})\n")
                f.write(f"  Pos: left={shp.left}, top={shp.top}, width={shp.width}, height={shp.height}\n")
                if shp.has_text_frame:
                    f.write(f"  Text:\n")
                    for p in shp.text_frame.paragraphs:
                        text = "".join(r.text for r in p.runs)
                        if not text:
                            text = p.text
                        f.write(f"    - {text}\n")
                f.write("\n")

dump_full(r'D:\AICTE\demoppt\TCE_Yquanto_SIH26263.pptx', r'D:\AICTE\dump_tce_yquanto.txt')
dump_full(r'D:\AICTE\demoppt\PRAMAAN-AICTE_SIH2026_Idea_Presentation.pptx', r'D:\AICTE\dump_pramaan_old.txt')
print("Dumped both full presentations successfully.")
