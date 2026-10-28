from pptx import Presentation

prs = Presentation(r"D:\AICTE\finaldemoppt\SIH2026-IDEA-Presentation-Format.pptx")
print(f"Slide count: {len(prs.slides)}")

for idx, slide in enumerate(prs.slides):
    print(f"\n--- Slide {idx+1} ---")
    for s in slide.shapes:
        txt = ""
        if s.has_text_frame:
            txt = s.text_frame.text.replace("\n", " ")[:50]
        print(f"  Shape: name='{s.name}', id={s.shape_id}, type={s.shape_type}, left={s.left/914400:.2f}in, top={s.top/914400:.2f}in, w={s.width/914400:.2f}in, h={s.height/914400:.2f}in, text='{txt}'")

